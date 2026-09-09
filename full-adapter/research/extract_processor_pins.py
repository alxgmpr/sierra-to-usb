"""Cross-check the RK3568 workbook against the exact RK3568B2 V2.3 PDF.

Run with full-adapter/.venv/bin/python from any directory. Never substitutes a
generic rectangular ball grid for the populated-ball list.
"""
from pathlib import Path
import hashlib
import json
import re

import openpyxl
import pymupdf

ROOT = Path(__file__).resolve().parents[1]
REF = ROOT / "reference"
BALL = re.compile(r"[A-Z]{1,2}[0-9]+")


def normalized(value):
    # PDF underscores are separate glyphs and sometimes appear on a new line.
    return re.sub(r"[^A-Z0-9/]", "", value.upper())


def main():
    workbook_path = REF / "RK3568_PinOut_V1.1_20210202.xlsx"
    pdf_path = REF / "rk3568b2-v2.3.pdf"
    workbook = openpyxl.load_workbook(workbook_path, data_only=True)
    sheet = workbook["05.PinOut List"]
    # Resolve merged metadata cells explicitly; do not forward-fill across groups.
    merged = {}
    for cell_range in sheet.merged_cells.ranges:
        value = sheet.cell(cell_range.min_row, cell_range.min_col).value
        for row in range(cell_range.min_row, cell_range.max_row + 1):
            for col in range(cell_range.min_col, cell_range.max_col + 1):
                merged[row, col] = value
    pins = {}
    for index, cells in enumerate(sheet.values, 1):
        row = [merged.get((index, col), value) for col, value in enumerate(cells, 1)]
        if not isinstance(row[0], str) or not BALL.fullmatch(row[0]):
            continue
        assert row[0] not in pins
        pins[row[0]] = {
            "ball": row[0], "workbook_name": row[1],
            "power_domain": row[2], "workbook_voltage": row[3],
            "pin_type": row[4], "reset_type": row[5], "reset_pull": row[6],
            "functions": [v for v in row[12:18] if v],
            "reference_net": row[18], "description": row[21],
            "workbook_row": index,
        }
    assert len(pins) == 636
    grid = {}
    for row_index, row in enumerate(list(workbook["04.Ball Map"].values)[1:29]):
        assert isinstance(row[0], str)
        for column, value in enumerate(row[1:29], 1):
            if value in (None, "NP"):
                continue
            ball = f"{row[0]}{column}"
            assert ball in pins
            assert normalized(value) == normalized(pins[ball]["workbook_name"])
            grid[ball] = [round((column - 14.5) * .65, 6),
                          round((row_index - 13.5) * .65, 6)]
    assert set(grid) == set(pins)
    pdf = pymupdf.open(pdf_path)
    extracted = {}
    for page_index in range(22, 29):
        for table in pdf[page_index].find_tables().tables:
            for row in table.extract():
                values = [v for v in row if v]
                for index, value in enumerate(values):
                    if value in pins and index and not BALL.fullmatch(values[index - 1]):
                        assert value not in extracted
                        extracted[value] = (values[index - 1], page_index + 1)
    assert set(extracted) == set(pins)
    differences = []
    for ball, pin in pins.items():
        name, page = extracted[ball]
        pin.update(xy_mm=grid[ball], b2_pdf_cell=name, b2_pdf_page=page)
        match = normalized(name) == normalized(pin["workbook_name"])
        pin["name_comparison"] = "MATCH" if match else "REVIEW_REQUIRED"
        pin["canonical_name"] = pin["workbook_name"]
        if not match and ball == "H24":
            # Table 2-1 is visibly truncated at the page break. Table 2-3 p34
            # contains the complete name, visually reviewed against the workbook.
            pin["name_comparison"] = "MATCH_TO_B2_TABLE_2_3_PAGE_34"
        elif not match and ball in {"AC14", "AC17", "AD12", "AD14", "AD15", "AD17", "AD18", "AE12", "AE15", "AE18"}:
            base = pin["workbook_name"]
            alias = base.replace("MIPI_DSI_", "LVDS_")
            assert normalized(name) == normalized(base + "/" + alias)
            pin["canonical_name"] = base + "/" + alias
            pin["name_comparison"] = "B2_ADDS_LVDS_ALIAS_VISUALLY_REVIEWED"
        elif not match and ball in {"W15", "Y15"}:
            rail = "0V9" if ball == "W15" else "1V8"
            pin["canonical_name"] = f"MIPI_DSI_TX1/LVDS_TX1_AVDD_{rail}"
            pin["name_comparison"] = "B2_ADDS_LVDS_POWER_ALIAS_VISUALLY_REVIEWED"
        if not match:
            differences.append(ball)
        ddr3 = [v for v in pin["workbook_name"].split("/") if v.startswith("DDR3_")]
        if ddr3:
            assert len(ddr3) == 1
            pin["ddr3_function"] = ddr3[0]
    payload = {
        "status": "BALL_POPULATION_AND_PIN_NAMES_CROSS_CHECKED",
        "sources": {p.name: hashlib.sha256(p.read_bytes()).hexdigest()
                    for p in (workbook_path, pdf_path)},
        "view": "PCB top view, A1 upper left; package bottom-view drawing mirrored",
        "pitch_mm": .65, "populated_balls": len(pins),
        "name_matches": len(pins) - len(differences),
        "name_difference_balls": sorted(differences),
        "pins": sorted(pins.values(), key=lambda p: (p["xy_mm"][1], p["xy_mm"][0])),
    }
    assert all(p["name_comparison"] != "REVIEW_REQUIRED" for p in pins.values())
    path = ROOT / "research" / "processor-pins.json"
    path.write_text(json.dumps(payload, indent=2, ensure_ascii=False) + "\n")
    print(json.dumps({k: v for k, v in payload.items() if k != "pins"}, indent=2))


if __name__ == "__main__":
    main()
